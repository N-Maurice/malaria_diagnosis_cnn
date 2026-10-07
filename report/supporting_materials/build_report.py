"""Build the evidence-based Word report from the fixed supporting snapshots.

Run from the repository root using .venv/bin/python. LibreOffice updates the
Word fields and produces the verified final Word/PDF artifacts afterward.
"""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path.cwd(); REPORT=ROOT/'report'; FIG=REPORT/'figures'; SUP=REPORT/'supporting_materials'
david=pd.read_csv(SUP/'david_validation_record.csv')
sarah=pd.read_csv(SUP/'sarah_validation_record.csv')
history=pd.read_csv(SUP/'maurice_baseline_history_snapshot.csv')
maurice_cfg=json.loads((SUP/'maurice_baseline_config_snapshot.json').read_text())
test=json.loads((SUP/'test_metrics_from_notebook_counts.json').read_text())
N=len(history)
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.55);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.3)
styles=doc.styles
normal=styles['Normal'];normal.font.name='Liberation Sans';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(3);normal.paragraph_format.line_spacing=1.08
for name,size,color in [('Heading 1',15,'17364D'),('Heading 2',11.5,'1E6B79'),('Heading 3',10.5,'17364D')]:
 s=styles[name];s.font.name='Liberation Sans';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color)
 s.paragraph_format.space_before=Pt(6);s.paragraph_format.space_after=Pt(3)
 s.paragraph_format.keep_with_next=True
styles['Caption'].font.name='Liberation Sans';styles['Caption'].font.size=Pt(9)
styles['Caption'].font.color.rgb=RGBColor.from_string('455A64')
styles['Caption'].paragraph_format.space_after=Pt(6)
for name in ['TOC 1','TOC 2','TOC 3']:
 if name not in styles: styles.add_style(name,1)
 styles[name].font.name='Liberation Sans';styles[name].font.size=Pt(9)
 styles[name].paragraph_format.space_after=Pt(1)
 styles[name].paragraph_format.line_spacing=1.0
for style in styles:
 if style.type==1 and style.element.rPr is not None:
  fonts=style.element.rPr.find(qn('w:rFonts'))
  if fonts is not None:
   for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
    fonts.attrib.pop(qn('w:'+attr),None)
sec.different_first_page_header_footer=True
hp=sec.header.paragraphs[0];hp.text='MALARIA CELL CLASSIFICATION  |  FORMATIVE 2';hp.style='Caption'
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Evidence snapshot: 7 October 2026  •  Page ').font.size=Pt(8)

def field(p,instruction,text=''):
 r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
 r=p.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=instruction;r._r.append(i)
 r=p.add_run();s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');r._r.append(s)
 if text: p.add_run(text)
 r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)
field(fp,' PAGE ')
settings=doc.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
figures=[];tables=[]
def p(text,style=None): return doc.add_paragraph(text,style)
PAGE_PENDING=False
def h(text,level=1):
 global PAGE_PENDING
 para=doc.add_heading(text,level)
 if PAGE_PENDING:
  para.paragraph_format.page_break_before=True
  PAGE_PENDING=False
 return para
def page():
 global PAGE_PENDING
 PAGE_PENDING=True
def cap(kind,title):
 seq=figures if kind=='Figure' else tables;seq.append(title)
 para=doc.add_paragraph(style='Caption');para.add_run(kind+' ');field(para,f' SEQ {kind} \\* ARABIC ',str(len(seq)));para.add_run('. '+title)
 para.paragraph_format.keep_with_next=kind=='Table'
 return len(seq)
def image(filename,title,width=6.85):
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(2);para.paragraph_format.keep_with_next=True
 para.add_run().add_picture(str(FIG/filename),width=Inches(width))
 return cap('Figure',title)
def table(title,headers,rows,widths=None,font=9):
 cap('Table',title)
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.autofit=False
 if widths:
  for column,width in zip(t.columns,widths): column.width=Inches(width)
 for cell,text in zip(t.rows[0].cells,headers): cell.text=text
 trPr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trPr.append(rep)
 for row in rows:
  for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2)
    para.paragraph_format.line_spacing=1.02
    for run in para.runs:run.font.size=Pt(font)
 for c in t.rows[0].cells:
  for r in c.paragraphs[0].runs:r.bold=True
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

def placeholder(text):
 para=p('[PLACEHOLDER – '+text+']');para.paragraph_format.space_before=Pt(6)
 for r in para.runs:r.italic=True;r.font.color.rgb=RGBColor.from_string('90531E')

def pct(x):return f'{x*100:.2f}%'
def namebreak(s):return s.replace('_','_\u200b')

# 1. Title page
p('AFRICAN LEADERSHIP UNIVERSITY','Subtitle')
p('Introduction to Machine Learning • Formative 2','Subtitle')
for _ in range(3):p('')
p('Malaria Cell Image\nClassification','Title')
p('Final Project Report','Subtitle')
p('Custom ResNet • EfficientNetB0 • Additional Custom CNN')
p('David • Sarah • Laura • Maurice')
p('7 October 2026')
p('Evidence-based account of the current implementation and recorded results.','Caption')
p('Full names, student IDs, lecturer name and group number are not supplied in the project. Add these details before submission.','Caption')
p('Scope: completed results, partial training, planned experiments and missing work are distinguished throughout. No missing results have been invented.','Caption')

# 2. Navigation
page();h('Table of Contents')
field(p(''),' TOC \\o "1-2" \\h \\z \\u ','Update fields to display contents.')

# 3. Caption lists
page();h('List of Figures')
field(p(''),' TOC \\h \\z \\c "Figure" ','Update fields to display figures.')
h('List of Tables')
field(p(''),' TOC \\h \\z \\c "Table" ','Update fields to display tables.')
p('Reading the evidence','Heading 2')
p('Recorded values come from preserved notebooks or experiment records. Derived metrics are calculated from explicit confusion counts or rounded precision/recall. Partial histories are not final results. Planned settings have no measured outcome. Missing implementations and artifacts remain labelled.')

# 4. Summary, contributions, introduction and literature
page();h('1. Summary, Contributions and Introduction')
h('1.1 Executive summary',2)
p(f'The project classifies cropped microscopic cells as Uninfected (0) or Parasitized (1). David records eight custom ResNet experiments; Sarah records seven EfficientNetB0 experiments. Their saved final tests exceed the assignment’s respective sensitivity/specificity targets. Maurice has an implemented plain CNN and {N} recorded baseline epochs, but no final evaluation. Laura’s assigned transfer model remains a stub. Different splits, conflicting Sarah records and patient/slide overlap limit comparison and generalization.')
h('1.2 Evidence-backed individual contributions',2)
table('Team ownership, contributions and completion state',['Member','Observed work','Status'],[
 ['David','Shared split manifests; subclassed ResNet; runner and Colab notebook; eight experiment records; final evaluation and error analysis.','Implemented; recorded results.'],
 ['Sarah','EfficientNetB0; Kaggle notebook; seven run records; evaluation, error and Grad-CAM figures; literature draft.','Implemented; records require reconciliation.'],
 ['Maurice',f'Plain CNN; runner; sixteen-section notebook; eight planned settings; {N}-epoch history, checkpoint and local logs.','Implemented; partial training.'],
 ['Laura','Assigned Transfer Learning Model 2 in README; model module contains a TODO.','Ownership confirmed; implementation not evidenced.']],widths=[.65,4.5,1.7],font=8.5)
p('Sources: README ownership, docs/contribution_log.md and member files. This contribution sheet does not certify unrecorded work or equal effort.','Caption')
h('1.3 Problem, objectives and related work',2)
p('Malaria’s disproportionate burden in Africa makes reliable diagnostic support relevant to Sub-Saharan healthcare [1]. This project addresses cell-image classification, not patient diagnosis, species identification or parasite-density measurement. Its objectives are one model per member, at least seven meaningful experiments per model, traceable TensorBoard evidence, validation-based selection and held-out evaluation. Targets are ≥90% sensitivity/specificity for custom models and ≥95% sensitivity/≥90% specificity for transfer models; they are assignment targets, not clinical certification [P1].')
p('Rajaraman et al. compare custom and pretrained malaria classifiers and consider patient-level evaluation [2]. He et al.’s residual learning motivates David’s shortcuts [3]; EfficientNet’s balanced scaling motivates Sarah’s compact pretrained backbone [4]. Maurice’s explicit plain CNN provides a distinct hand-built design. Batch normalization and dropout motivate controlled optimization/regularization comparisons [5,6]. Medical imaging reviews emphasize evaluation challenges [7]; Grad-CAM supports inspection of class-related regions but cannot prove correct medical reasoning [8]. These sources justify the designs and the need for cautious evaluation rather than accuracy alone.')

# 5. Methodology and design
page();h('2. Methodology, Design and Implementation')
h('2.1 Dataset and split integrity',2)
p('The project uses the NIH/NLM malaria cell dataset described by the assignment and repository: 27,558 local images, with 13,779 per class. Images are already cell crops; no segmentation is implemented. Folder names provide labels. Local paths and labels match the committed manifests, and split filenames are disjoint. Source authenticity and extraction checksums are not preserved [P2].')
table('Committed shared manifest composition',['Split','Uninfected','Parasitized','Total','Use'],[
 ['Train','9,645','9,645','19,290','Fit weights'],['Validation','2,067','2,067','4,134','Select settings/checkpoints'],['Test','2,067','2,067','4,134','Evaluate selected model']],widths=[.8,1.05,1.05,.8,3.15],font=9)
p('Shared splitting is stratified 70/15/15 with seed 42. Sarah instead shuffles and slices an independent dataframe; her test supports are 2,044 Uninfected and 2,090 Parasitized. Equal total size does not establish shared split identity or stratification. All 148 parsed test patient/slide-like filename codes occur in shared training; 555 test filenames are unparsed. This heuristic cannot establish unseen-patient evaluation.')
h('2.2 Preprocessing, architecture and workflow',2)
image('architecture_overview.png','Implemented architecture paths; Laura’s second transfer architecture is not yet evidenced.',width=6.2)
p('Shared utilities decode RGB, resize custom-model inputs to 128×128, normalize to [0,1], batch and prefetch. Only training is shuffled/augmented; ordered validation/test predictions align with manifests. Sarah uses 224×224 and EfficientNet’s input contract. David’s Keras Layer/Model subclasses use two-convolution residual blocks, identity/projection shortcuts, global average pooling and sigmoid classification. His final record has four stages, two blocks each and 2,800,097 parameters. Sarah replaces EfficientNetB0’s ImageNet head with pooling, dropout and sigmoid; her selected record unfreezes the last 40 layers. Maurice’s three-block baseline has two convolutions per block and 287,137 parameters [P3–P6].')
p('The workflow is data → manifests → streamed training/validation → validation-based selection → held-out evaluation → error analysis/Grad-CAM. Shared code is in src/, models in members/, notebooks in notebooks/ and manifests in splits/. Weights/results/logs are ignored artifacts. Custom runners initialize fresh models, save settings/history and restore best validation-loss weights. Maurice additionally freezes split/checkpoint hashes and threshold before cached test evaluation. Notebooks are the working interface; no deployed application is evidenced.')
h('2.3 Evaluation protocol',2)
p('Accuracy = (TP+TN)/N; precision = TP/(TP+FP); sensitivity/recall = TP/(TP+FN); specificity = TN/(TN+FP). F1 is the harmonic mean of precision and recall. ROC-AUC summarizes discrimination across thresholds. Threshold 0.5 classifies probabilities ≥0.5 as infected. Test data must not guide settings or threshold selection. Streaming training AUC and exact probability-based ROC-AUC are distinguished.')

# 6. David
page();h('3. Experiments: David’s Custom ResNet')
h('3.1 Eight recorded validation comparisons',2)
p('Common settings: shared manifests, 128×128 RGB, batch size 32, BatchNorm, binary cross entropy, fresh initialization and best validation-loss weights. A/P/R/F1 are percentages; R is infected-class sensitivity. AUC is on a 0–1 scale. Exact recorded names are retained [P4].')
notes=['Reference; best 6/11 epochs.','More capacity; best 6/11.','No skips; best 20/20.','Lower AUC; best 1/6.','Lower recall; best 7/12.','Recovery; best 20/20.','Highest recall; 20/20.','Selected; best 30/30.']
configs=['Adam 1e-3; 3 stages, 1 block/stage.','Adam 1e-3; 4 stages, 2 blocks/stage.','Deep layout; skip paths removed.','02 + flips, turns, colour jitter.','04 + dropout .3 and L2 1e-4.','05; Adam LR 3e-4.','06; SGD/Nesterov, LR 1e-2.','From 06; plateau LR reduction; 30 epochs.']
rows=[]
for i,r in david.iterrows():
 m=f'A {pct(r.accuracy)} / P {pct(r.precision)}\nR {pct(r.recall)} / F1 {pct(r.f1)}\nAUC {r.roc_auc:.4f}'
 rows.append([namebreak(r.run_name),configs[i],m,notes[i]])
table('ResNet experiment progression from recorded validation results',['Exact run name','Configuration / change','Validation metrics','Outcome'],rows,widths=[1.95,1.65,1.95,1.3],font=8.5)
p('Specificity for runs 01–08: 97.92%, 97.34%, 97.10%, 97.10%, 98.02%, 97.73%, 96.47% and 98.06%. The complete supporting CSV preserves values and run times. The source records commit eced63e and combined manifest digest prefix c5f90e8c1a71. Raw runs are described on David’s Drive, but a viewable event-log link is missing.')
h('3.2 Interpretation and selection',2)
p('Depth modestly improves the baseline, but removing shortcuts also reaches a slightly higher accuracy at a longer training cap. This does not prove that skips improve final accuracy in this comparison. Augmentation and dropout/L2 initially reduce recall and AUC: useful failed hypotheses. Smaller learning rate later improves the record. Run 08 has the best recorded accuracy/AUC, while 07 has higher sensitivity and lower specificity. Selection therefore involves a trade-off. Run 08’s best epoch is its last, leaving convergence uncertain. Regularized losses include L2 and are not directly comparable with unregularized classification loss.')

# 7. Sarah
page();h('4. Experiments: Sarah’s EfficientNetB0')
h('4.1 Seven recorded fine-tuning comparisons',2)
p('Input is 224×224 with training augmentation. Defaults are Adam, batch 32 and dropout .3; later runs allow 15 epochs with early stopping. Table values are JSON/Markdown records. F1 is derived from rounded precision/recall and is approximate; per-run specificity is missing [P5].')
configs=['Frozen base; Adam 1e-3; 5 epochs.','Last 20 layers; Adam 1e-5; 5 epochs.','Last 20 layers; 15-epoch cap; early stop.','SGD momentum .9; LR 1e-3; last 20.','Dropout .5; Adam 1e-5; last 20.','Last 40 layers; Adam 1e-5; dropout .3.','From 06; batch size 64.']
notes=['Frozen baseline.','Recall increases.','Recall exceeds 95%.','Similar recorded scores.','Small recorded change; conflict.','Selected candidate.','Similar scores; conflict.']
rows=[]
for i,r in sarah.iterrows():
 m=f'A {pct(r.val_accuracy)} / P {pct(r.val_precision)}\nR {pct(r.val_recall_sensitivity)} / F1 ≈{pct(r.derived_f1)}\nAUC {r.val_auc:.4f}'
 rows.append([namebreak(r.run_name),configs[i],m,notes[i]])
table('EfficientNet progression from recorded validation summaries',['Exact run name','Configuration / change','Validation metrics','Outcome'],rows,widths=[1.95,1.65,1.95,1.3],font=8.5)
h('4.2 Interpretation and unresolved records',2)
p('Recorded recall rises from 93.18% in the frozen baseline to 96.21% in run 06; AUC rises from .9824 to .9934. Dropout, optimizer and batch changes have small recorded effects. These are within-owner comparisons on Sarah’s independent split, not proof of superiority on a common test set.')
p('Run 05’s visible final-epoch accuracy is about 96.06%, versus a 96.25% summary; run 07 shows about 96.30%, versus 96.40%. Record cells use manually entered values. No validation prediction CSV resolves checkpoint differences, so the table remains labelled as recorded summaries. Saved notebook setup is incomplete, visible fit calls lack TensorBoard callbacks and local raw event logs are absent. The final test is reported consistently from its explicit notebook confusion matrix rather than conflicting summaries/standalone PNGs.')

# 8. Maurice/Laura/testing
page();h('5. Maurice, Laura and Verification')
h('5.1 Partial baseline and remaining experiment plan',2)
p(f'Maurice’s baseline uses Adam 1e-3, batch 32, three blocks and a 25-epoch cap, with batch normalization, dropout, L2 and augmentation off. The fixed snapshot has {N} epoch rows, a checkpoint, configuration, package snapshot and TensorBoard files. It has no completed metrics.json, validation predictions or frozen final test selection [P6].')
image('maurice_partial_training.png',f'Maurice’s {N} recorded baseline epochs: partial training, not final evaluation.',width=6.2)
last=history.iloc[-1]
p(f'Last recorded epoch: train accuracy {pct(last.accuracy)}, validation accuracy {pct(last.val_accuracy)}, validation loss {last.val_loss:.4f}, streaming AUC {last.val_auc:.4f}. These are epoch values, not final restored-checkpoint results.')
table('Maurice’s experiment plan and evidence state',['Exact run suffix¹','Change','Evidence'],[
 ['exp_01_baseline','Unregularized three-block CNN',f'{N} recorded epochs; partial'],['exp_02_batchnorm','Add batch normalization','Planned'],['exp_03_dropout_l2','Add dropout .3 and L2 1e-4','Planned'],['exp_04_augmentation','03 + flip, rotation and zoom','Planned'],['exp_05_lr_3e-4','03 + lower Adam learning rate','Planned'],['exp_06_sgd_momentum','03 + SGD/Nesterov, LR 1e-2','Planned'],['exp_07_batch64','03 + batch 64','Planned'],['exp_08_wider_aug_plateau','Fourth block + augmentation + LR schedule','Planned combined candidate']],widths=[1.8,3.4,1.65],font=8.5)
p('¹ Prefix each suffix with additional_custom_cnn_. Plans have no measured metrics. Combined changes cannot isolate one variable’s effect. The snapshot remains fixed even if training continues.','Caption')
h('5.2 Laura and synthetic verification',2)
p('Laura’s second transfer model is assigned but contains only a TODO; no architecture, individual notebook or results are evidenced [P7]. All 11 synthetic tests passed during report preparation. They cover reproducible splits, metrics, pipelines, TensorBoard, connected Grad-CAM, CNN inference/checkpoint reload and frozen/cached evaluation. They verify code behavior, not full training reproduction or clinical accuracy [P8].')

# 9. Evaluation panels
page();h('6. Final Evaluation Visualizations')
h('6.1 David: selected ResNet run 08',2)
image('david_final_four_panel_row.png','ResNet: saved accuracy, loss, test confusion matrix and ROC curve together.')
p('Across 30 epochs, accuracy improves and loss falls, with validation dips around epochs 10–11 and 16 that recover. The best-epoch train/validation accuracies are 97.35%/97.24%, a .11-point gap. Validation instability is visible; the small gap does not prove generalization. The best checkpoint is at the epoch cap. Test counts are TN 2,020, FP 47, FN 89 and TP 1,978. Numerical ROC-AUC is .9946; the plot rounds it to about .995.')
h('6.2 Sarah: selected EfficientNet run 06',2)
image('sarah_final_four_panel_row.png','EfficientNet: saved notebook plots arranged in the required four-panel row.')
p('Train/validation accuracy improves and both losses decline over 15 epochs. Validation exceeds training slightly under different dropout/augmentation conditions; this does not prove absence of overfitting. The best-loss checkpoint’s relation to reported last-epoch metrics needs confirmation. The saved notebook shows TN 1,979, FP 65, FN 97 and TP 1,993; ROC-AUC .9925. The separate PNG instead shows FP 66, FN 95 and AUC .9928. This report consistently uses the notebook, not a mixture of artifacts.')
h('6.3 Missing final model evidence',2)
placeholder('Maurice: insert all four final plots in one row after selecting on validation and freezing held-out evaluation.')
placeholder('Laura: insert all four final plots after implementing and evaluating the second transfer model.')
p('Partial training curves cannot replace final evaluation. No unsupported test curve or matrix was generated to fill these gaps. High-resolution originals remain in report/figures/.')

# 10. Results and errors
page();h('7. Results, Comparison and Error Analysis')
h('7.1 Supported final tests',2)
rows=[]
for owner,model in [('David','ResNet, run 08'),('Sarah','EfficientNet, run 06')]:
 r=test[owner];rows.append([model,pct(r['accuracy']),pct(r['precision']),pct(r['recall']),pct(r['specificity']),pct(r['f1']),f"{r['roc_auc']:.4f}"])
rows.extend([['Maurice','Missing','Missing','Missing','Missing','Missing','Missing'],['Laura','Missing','Missing','Missing','Missing','Missing','Missing']])
table('Final test rates derived from saved notebook confusion counts',['Model','Accuracy','Precision','Recall / sensitivity','Specificity','F1','AUC'],rows,widths=[1.35,.85,.85,1.15,.9,.85,.9],font=8.5)
p('David meets the custom-model ≥90%/≥90% sensitivity/specificity target; Sarah meets the transfer-model ≥95%/≥90% target. The ResNet’s recorded accuracy is .63 percentage points higher, but splits, resolution, training and preprocessing differ. No repeated-seed variance, matched-prediction test or external evaluation supports a controlled ranking. Maurice and Laura cannot be ranked.')
h('7.2 Representative mistakes and likely causes',2)
table('David’s recorded examples at threshold 0.5',['Image identifier²','True → prediction','P(infected)','Possible cause'],[
 ['C75P36_ThinF_IMG_20150815_163707_cell_33.png','Uninfected → Parasitized','.9989','Stained inclusion; artifact or label ambiguity.'],['C128P89ThinF_IMG_20151004_131753_cell_99.png','Uninfected → Parasitized','.9983','Edge staining; confident artifact-based prediction possible.'],['C100P61ThinF_IMG_20150918_144348_cell_142.png','Parasitized → Uninfected','.0035','No obvious inclusion; faint/cropped signal or label ambiguity.']],widths=[2.35,1.3,.75,2.45],font=8.5)
p('² First two in Uninfected/, third in Parasitized/ under data/cell_images/. These are hypotheses, not corrected labels.','Caption')
image('sarah_error_examples_notebook.png','Sarah’s saved false-negative and false-positive examples with probabilities.',width=4.5)
p('David records 136 errors: 89 misses and 47 false alarms. Sarah’s notebook records 162: 97 misses and 65 false alarms. Her gallery supplies labels/probabilities but not filenames. Faint staining and dark debris are reported hypotheses without expert annotations. David’s final curves show transient instability; Sarah’s curves stay close under different training conditions. Augmentation, regularization and optimizer changes have mixed effects. Maurice’s short history cannot establish convergence; Laura has no curves or error evidence.')

# 11. Explainability
page();h('8. Model Explainability')
h('8.1 Correct and incorrect examples',2)
p('Grad-CAM uses gradients of a class score to weight convolutional feature maps and highlight positive attribution [8]. David exposes connected features/logits and targets the predicted class: infected logit or its negative for uninfected. Maurice exposes a connected last_conv/logit path but has no completed heatmap evidence.')
image('sarah_gradcam_notebook.png','Sarah’s correct infected, correct uninfected and incorrect positive examples.')
p('Sarah’s infected example highlights a dark structure near the boundary; the uninfected map includes edge/background-adjacent activity; the error emphasizes a dark inclusion. Her saved code targets infected probability even for the uninfected example, so that panel is not a negative-class explanation.')
image('david_gradcam_compact.png','David’s saved correct-infected, correct-uninfected and incorrect examples; originals, maps and overlays preserved.')
h('8.2 Interpretation, comparison and confidence',2)
p('David’s infected maps often emphasize localized staining; negative-class maps can be diffuse. Confident false positives also focus on inclusions in labelled uninfected cells. Apparent map differences cannot be attributed solely to custom versus pretrained architecture because inputs, target classes and map resolutions differ. Matched examples and a common protocol are needed. Edge attention and confident errors raise shortcut/artifact concerns rather than proving parasite localization. Expert review is missing. [PLACEHOLDER – add Maurice’s and Laura’s correct-infected, correct-uninfected and incorrect examples, with interpretation, after evaluation.]')

# 12. Discussion and conclusion
page();h('9. Challenges, Limitations and Conclusions')
h('9.1 Practical challenges and remaining limitations',2)
p('The initial Anaconda kernel lacked TensorFlow/TensorBoard. A project-local Python 3.13 environment was installed; imports and 11 synthetic tests passed. Empty dataset folders were resolved by copying existing local images and checking them against shared manifests. No GPU was detected locally. These setup fixes do not verify data provenance or reproduce completed remote training.')
p('Image-level disjointness does not establish patient independence: all 148 parsed shared test codes also appear in training. Duplicate-image content, external validation, calibration and uncertainty are not assessed. One reported seed per setting leaves variability unknown. Sarah’s independent split and conflicting summaries prevent a controlled cross-model comparison. Her raw TensorBoard evidence is absent locally; David’s described Drive path lacks a viewable link. Maurice’s evidence is partial and Laura’s implementation is missing. The four-member submission is therefore incomplete, even though the documented models exceed assignment targets.')
h('9.2 Conclusion',2)
p('The team has developed reusable data/evaluation infrastructure, a hand-built residual model, an EfficientNetB0 transfer model and an implemented plain CNN. Eight ResNet and seven EfficientNet experiment records support useful within-owner comparisons. Saved final tests show high cell-level discrimination, but remaining errors, split limitations and incomplete evidence restrict the conclusion. There is no established four-model ranking or clinical diagnostic validation. Missing contributions/results are disclosed rather than filled by assumptions.')
h('9.3 Recommendations and submission priorities',2)
p('Finish Maurice’s meaningful comparisons and Laura’s distinct transfer architecture; select on validation before final test evaluation. Reconcile Sarah’s exact split, checkpoint and artifact values. Provide every member’s viewable notebook, raw event-directory link and TensorBoard comparison exports. Fill institutional details and missing figure placeholders with real evidence. These actions address immediate assignment gaps.')
p('For stronger evaluation, obtain authoritative patient/slide groups, audit duplicates and test group-aware or external generalization. Repeat selected settings across seeds, quantify uncertainty and consider validation-only threshold selection/calibration. Expert review should test error and heatmap hypotheses. WHO’s regulatory considerations provide broader context for risk–benefit assessment and monitoring [9]; this academic prototype has no regulatory clearance. Each member must prepare the required on-camera defense and verify their contribution. The assignment’s below-20% AI-generated narrative requirement still needs original student review and revision; this document cannot certify compliance [P1].')

# 13. References
page();h('References')
p('Numbered citations are consistent throughout. External sources were checked on 7 October 2026; project evidence IDs are indexed in Appendix A.')
refs=[
('[1]','World Health Organization. (2025, December 4). Malaria. https://www.who.int/news-room/fact-sheets/detail/malaria'),
('[2]','Rajaraman, S., Antani, S. K., Poostchi, M., Silamut, K., Hossain, M. A., Maude, R. J., Jaeger, S., & Thoma, G. R. (2018). Pre-trained convolutional neural networks as feature extractors toward improved malaria parasite detection in thin blood smear images. PeerJ, 6, e4568. https://doi.org/10.7717/peerj.4568'),
('[3]','He, K., Zhang, X., Ren, S., & Sun, J. (2016). Deep residual learning for image recognition. CVPR, 770–778. https://arxiv.org/abs/1512.03385'),
('[4]','Tan, M., & Le, Q. V. (2019). EfficientNet: Rethinking model scaling for convolutional neural networks. PMLR, 97, 6105–6114. https://proceedings.mlr.press/v97/tan19a.html'),
('[5]','Ioffe, S., & Szegedy, C. (2015). Batch normalization: Accelerating deep network training by reducing internal covariate shift. PMLR, 37, 448–456. https://proceedings.mlr.press/v37/ioffe15.html'),
('[6]','Srivastava, N., Hinton, G., Krizhevsky, A., Sutskever, I., & Salakhutdinov, R. (2014). Dropout: A simple way to prevent neural networks from overfitting. JMLR, 15, 1929–1958. https://www.jmlr.org/papers/v15/srivastava14a.html'),
('[7]','Litjens, G., Kooi, T., Bejnordi, B. E., Setio, A. A. A., Ciompi, F., Ghafoorian, M., van der Laak, J. A. W. M., van Ginneken, B., & Sánchez, C. I. (2017). A survey on deep learning in medical image analysis. Medical Image Analysis, 42, 60–88. https://doi.org/10.1016/j.media.2017.07.005'),
('[8]','Selvaraju, R. R., Cogswell, M., Das, A., Vedantam, R., Parikh, D., & Batra, D. (2017). Grad-CAM: Visual explanations from deep networks via gradient-based localization. ICCV, 618–626. https://arxiv.org/abs/1610.02391'),
('[9]','World Health Organization. (2023). Regulatory considerations on artificial intelligence for health. https://www.who.int/publications/i/item/9789240078871')]
for label,text in refs:
 para=p(label+' '+text);para.paragraph_format.space_after=Pt(7);para.paragraph_format.left_indent=Inches(.25);para.paragraph_format.first_line_indent=Inches(-.25)

# 14. Evidence and links
page();h('Appendix A. Evidence Sources and Submission Links')
h('A.1 Project evidence register',2)
table('Sources supporting implementation, results and contributions',['ID','Source','Evidence'],[
 ['P1','Assignment PDF; extracted text in supporting_materials/','Required sections, targets, plots, defense and narrative policy'],
 ['P2','README; docs/dataset_notes.md; splits/*.csv','Dataset assumptions and verified manifest counts'],
 ['P3','src/{config,data_utils,eval_utils,logging_utils,gradcam_utils}.py','Shared data, metrics, tracking and explanation contracts'],
 ['P4','members/david_custom_resnet/; notebooks/david/custom_resnet_malaria.ipynb','Residual architecture, eight records, saved test/error/Grad-CAM evidence'],
 ['P5','members/sarah_transfer_model_1/; notebooks/sarah/TheNoteBook.ipynb; existing report draft','Transfer architecture, seven records, saved outputs and discrepancies'],
 ['P6','members/maurice_custom_cnn/; Maurice notebook; fixed history/config snapshots','Plain architecture, partial run, planned settings and provenance'],
 ['P7','members/laura_transfer_model_2/model.py; workspace README','Assigned ownership; missing implementation'],
 ['P8','tests/; supporting_materials/verification_tests.txt; contribution log','Synthetic checks and documented contributions']],widths=[.45,3.55,2.85],font=8.5)
h('A.2 Required external links',2)
table('Notebook and raw TensorBoard evidence links still required',['Owner','Notebook evidence','Raw logs / comparison exports'],[
 ['David','Local notebook exists. [PLACEHOLDER – viewable Colab link]','Drive path recorded. [PLACEHOLDER – open-access raw-log link and comparisons]'],
 ['Sarah','Local notebook exists. [PLACEHOLDER – viewable notebook link]','[PLACEHOLDER – raw event directories and scalar comparisons]'],
 ['Maurice','Local notebook exists. [PLACEHOLDER – viewable notebook link]','Partial local logs. [PLACEHOLDER – complete runs and share raw logs]'],
 ['Laura','[PLACEHOLDER – implement and share individual notebook]','[PLACEHOLDER – complete tracking/runs and share evidence]']],widths=[.6,2.95,3.3],font=8.5)
p('The PDF requires open-access raw TensorBoard directories for all required runs, plus high-resolution run-comparison screenshots/exports. A filesystem path is not a viewable link. Oral defense requires each member on camera for approximately 10–15 minutes, including a two-minute overview and investigative questions; completion is not evidenced.')
p('Repository URL recorded in David’s notebook: https://github.com/N-Maurice/malaria_diagnosis_cnn (accessibility not verified). Course dataset-access notebook: https://drive.google.com/file/d/15CXWfcOND1he5ibp9PtH-716kJwRW3EW/view. This is not a model notebook or log evidence link.')

# 15. Audit
page();h('Appendix B. Conflicts and Requirements Audit')
h('B.1 Reporting decisions for conflicting evidence',2)
table('Material conflicts and evidence-based treatment',['Issue','Conflict','Treatment'],[
 ['Sarah split','Log calls split stratified/matched; code shuffles/slices; supports differ.','Use explicit notebook supports; no shared-split claim.'],
 ['Sarah test metrics','JSON: sensitivity .9545/spec .9677. Notebook counts: 97 FN/65 FP; rates .9536/.9682.','Derive rates from explicit counts.'],
 ['Sarah final PNG','Standalone: FP66/FN95/AUC .9928; notebook: FP65/FN97/AUC .9925.','Use notebook consistently; owner must confirm selected checkpoint.'],
 ['Sarah validation 05/07','Visible epoch metrics differ from manually entered summaries.','Label summaries as records; no exact reproduction claim.'],
 ['Stale scaffold / evolving training','Old README says no models; current modules exist. Maurice training is continuing.','Use implemented source and a fixed partial-run snapshot.']],widths=[1.2,3.1,2.55],font=8.5)
h('B.2 Required sections and remaining deliverables',2)
table('Assignment requirements mapped to this concise report',['Requirement','Coverage','Remaining evidence'],[
 ['Introduction, literature, objectives','Section 1; References','Student narrative revision'],['Dataset, preprocessing, architecture rationale','Section 2','Laura’s second architecture'],
 ['7+ meaningful runs/model; five metrics/run','Sections 3–5','Maurice incomplete; Laura absent; Sarah specificity missing'],
 ['Four final plots together and interpreted','Section 6','Maurice/Laura final panels'],['Comparison and representative errors','Section 7','Common split, exact Sarah error filenames, missing-owner examples'],
 ['Dedicated explainability and correct/incorrect cases','Section 8','Maurice/Laura maps; expert interpretation review'],['Critical discussion, limitations and conclusion','Section 9','External/group-aware evaluation remains future work'],
 ['Notebook/log links and tracking screenshots','Appendix A','Actual share links and TensorBoard exports'],['Contribution sheet and oral defense','Section 1; Appendix A','Institutional details, member confirmation and defense']],widths=[2.35,1.55,2.95],font=8.5)
p('Supporting files preserve the fixed history, experiment records, derived metrics, source hashes and 11-test verification output. Figures are saved notebook evidence or graphs/diagrams derived from actual records/code. No model was trained to manufacture missing report results. Replace placeholders only with completed, traceable evidence.')

# Document metadata and output
# Core metadata
doc.core_properties.title='Malaria Cell Image Classification — Final Project Report'
doc.core_properties.subject='Evidence-based Formative 2 report; current implementation and results'
doc.core_properties.author='Project team: David, Sarah, Laura and Maurice (AI-assisted assembly)'
doc.core_properties.keywords='malaria; CNN; ResNet; EfficientNet; Grad-CAM; evidence-based report'
doc.core_properties.comments='Generated from project evidence. Student review/revision required; no originality compliance is certified.'
path=REPORT/'Final_Project_Report.docx';doc.save(path)
(SUP/'caption_inventory.json').write_text(json.dumps({'figures':figures,'tables':tables},indent=2))
print('Saved',path,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'figures',len(figures))
